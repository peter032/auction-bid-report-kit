#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
네이버 부동산 매물 카드 원문 → 표(markdown) + 엑셀(.xlsx)

사용:
  python parse_listings.py <입력.json> [출력.xlsx] [--title "제목"]

입력: scrape.js가 뽑은 카드 원문 텍스트 배열(JSON)
      예) "오전동 복합상가월세 2,000/130상가점포2018.11(8년차)계약82.4㎡ (전용29.67)1/2층..."
출력: stdout = 마크다운 표(요청 컬럼) + 요약 / 지정 시 .xlsx 저장
"""
import sys, json, re, argparse, io

# Windows 콘솔 UTF-8 강제
try:
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
except Exception:
    pass

PY = 3.305785  # 1평 = 3.305785 ㎡


def maemae_to_man(s):
    """'2억 5,000' / '7억' / '5억 6,000' → 만원 정수"""
    s = s.replace(',', '').strip()
    m = re.match(r'(?:(\d+)억)?\s*(\d+)?', s)
    eok = int(m.group(1)) if m and m.group(1) else 0
    man = int(m.group(2)) if m and m.group(2) else 0
    return eok * 10000 + man


def parse_card(t):
    r = dict(typ='', floor='', jeon=None, py=None, trade=None,
             deposit=None, rent=None, maemae=None, maemae_txt=None,
             grouped=False, ncnt=1, raw=t)
    # 거래유형/가격
    m = re.search(r'월세\s*((?:\d+억)?[\s\d,]*)/([\d,]+)', t)
    if m:
        r['trade'] = '월세'
        r['deposit'] = maemae_to_man(m.group(1))
        r['rent'] = int(m.group(2).replace(',', ''))
    else:
        m = re.search(r'전세\s*((?:\d+억)?[\s\d,]*)', t)
        if m:
            r['trade'] = '전세'
            r['deposit'] = maemae_to_man(m.group(1))
        else:
            m = re.search(r'매매\s*((?:\d+억)?[\s\d,]*(?:~\s*(?:\d+억)?[\s\d,]*)?)', t)
            if m:
                r['trade'] = '매매'
                seg = m.group(1).strip()
                if '~' in seg:
                    r['maemae_txt'] = re.sub(r'\s+', ' ', seg)
                else:
                    r['maemae'] = maemae_to_man(seg)
    # 전용면적
    m = re.search(r'전용([\d.]+)\)', t)
    r['jeon'] = float(m.group(1)) if m else None
    # 층 (해당층/총층)
    m = re.search(r'\)(B?\d+)/(\d+)층', t) or re.search(r'(B?\d+)/(\d+)층', t)
    r['floor'] = (m.group(1) + '층') if m else None
    # 유형 (예: "오전동 복합상가" / "단지내상가" / "사무실")
    m = re.search(r'([가-힣]+동)?\s*(복합상가|단지내상가|상가|사무실|근린상가)', t)
    r['typ'] = m.group(2) if m else ''
    # 중개사 그룹
    r['grouped'] = '곳에서 등록' in t
    g = re.search(r'중개사\s*(\d+)곳', t)
    r['ncnt'] = int(g.group(1)) if g else 1
    # 파생
    if r['jeon']:
        r['py'] = r['jeon'] / PY
    if r['rent'] and r['py']:
        r['rent_ppp'] = round(r['rent'] / r['py'], 1)
    else:
        r['rent_ppp'] = None
    if r['maemae'] and r['py']:
        r['maemae_ppp'] = round(r['maemae'] / r['py'])
    else:
        r['maemae_ppp'] = None
    return r


def fmt(x, dec=2):
    if x is None:
        return ''
    if isinstance(x, float):
        return f"{x:,.{dec}f}"
    return f"{x:,}"


def print_table(rows):
    print("| # | 층 | 전용평수 | 임대보증금 | 월임대료 | 평당 월임대 단가 | 매매 금액 | 평당 매매 단가 | 비고 |")
    print("|---|---|---|---|---|---|---|---|---|")
    for i, r in enumerate(rows, 1):
        note_parts = [r['typ']]
        if r['grouped']:
            note_parts.append(f"중개사{r['ncnt']}곳")
        note = ' · '.join([p for p in note_parts if p])
        maemae_disp = fmt(r['maemae']) if r['maemae'] else (r['maemae_txt'] or '')
        print("| {} | {} | {} | {} | {} | {} | {} | {} | {} |".format(
            i, r['floor'] or '', (fmt(r['py']) + '평') if r['py'] else '',
            fmt(r['deposit']), fmt(r['rent']),
            fmt(r['rent_ppp'], 1), maemae_disp, fmt(r['maemae_ppp']), note))


def print_summary(rows):
    wolse = [r for r in rows if r['trade'] == '월세' and r['rent_ppp']]
    maemae = [r for r in rows if r['trade'] == '매매' and r['maemae_ppp']]
    print("\n### 요약")
    print(f"- 총 **{len(rows)}건** (월세 {sum(1 for r in rows if r['trade']=='월세')} · "
          f"매매 {sum(1 for r in rows if r['trade']=='매매')} · "
          f"전세 {sum(1 for r in rows if r['trade']=='전세')})")
    if wolse:
        ppps = sorted(r['rent_ppp'] for r in wolse)
        print(f"- 평당 월임대 단가: **{ppps[0]:.1f} ~ {ppps[-1]:.1f}만원/평**")
    if maemae:
        ppps = sorted(r['maemae_ppp'] for r in maemae)
        print(f"- 평당 매매 단가: **{ppps[0]:,} ~ {ppps[-1]:,}만원/평**")
    grp = sum(1 for r in rows if r['grouped'])
    if grp:
        print(f"- ※ 중개사 중복등록 그룹 {grp}건 포함 (네이버 '매물 N' 표기와 카드수 차이 사유)")


def write_xlsx(rows, path, title):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    wb = Workbook(); ws = wb.active; ws.title = '매물목록'
    ws.merge_cells('A1:I1'); c = ws['A1']; c.value = title
    c.font = Font(bold=True, size=13)
    ws['A2'] = '단위: 만원 | 전용평수=전용㎡÷3.3058 | 평당단가=금액÷전용평'
    ws['A2'].font = Font(size=9, color='666666')
    headers = ['#', '층', '전용평수', '임대보증금', '월임대료',
               '평당 월임대 단가', '매매 금액', '평당 매매 단가', '비고']
    hr = 4
    hf = PatternFill('solid', fgColor='2E7D32'); hfont = Font(bold=True, color='FFFFFF', size=10)
    thin = Side(style='thin', color='D0D0D0'); bd = Border(thin, thin, thin, thin)
    for j, h in enumerate(headers, 1):
        cell = ws.cell(hr, j, h); cell.fill = hf; cell.font = hfont; cell.border = bd
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    r = hr + 1
    for i, x in enumerate(rows, 1):
        note = ' · '.join([p for p in [x['typ'], (f"중개사 {x['ncnt']}곳" if x['grouped'] else '')] if p])
        maemae_val = x['maemae'] if x['maemae'] else (x['maemae_txt'] if x['trade'] == '매매' else None)
        vals = [i, x['floor'], round(x['py'], 2) if x['py'] else None,
                x['deposit'], x['rent'], x['rent_ppp'], maemae_val, x['maemae_ppp'], note]
        for j, v in enumerate(vals, 1):
            cell = ws.cell(r, j, v); cell.border = bd
            cell.alignment = Alignment(horizontal='center' if j in (1, 2) else 'right', vertical='center')
            if j in (4, 5, 7, 8) and isinstance(v, (int, float)):
                cell.number_format = '#,##0'
        if i % 2 == 0:
            for j in range(1, 10):
                ws.cell(r, j).fill = PatternFill('solid', fgColor='F5F9F5')
        r += 1
    for j, w in enumerate([4, 6, 10, 12, 10, 16, 16, 15, 20], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = 'A5'
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('input')
    ap.add_argument('xlsx', nargs='?', default=None)
    ap.add_argument('--title', default='네이버 부동산 매물 목록')
    a = ap.parse_args()

    data = json.load(open(a.input, encoding='utf-8'))
    rows = [parse_card(t) for t in data]

    print(f"# {a.title}\n")
    print_table(rows)
    print_summary(rows)

    if a.xlsx:
        write_xlsx(rows, a.xlsx, a.title)
        print(f"\n엑셀 저장: {a.xlsx}")


if __name__ == '__main__':
    main()
