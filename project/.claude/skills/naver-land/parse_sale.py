# -*- coding: utf-8 -*-
"""매매 매물 카드 원문 -> 층/전용평/평당단가/매매가/임차내용 표 + 엑셀 (중복 제거)"""
import sys, json, re, io, argparse
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

PY = 3.305785


def won(s):
    s = s.replace(',', '').strip()
    m = re.match(r'^(?:(\d+)억)?\s*(\d+)?$', s)
    if not m:
        return None
    eok = int(m.group(1)) if m.group(1) else 0
    man = int(m.group(2)) if m.group(2) else 0
    v = eok * 10000 + man
    return v or None


TENANT_KW = ['임차', '세입자', '세끼고', '세 끼고', '월세', '임대', '보증금', '수익률', '수익율',
             '렌트', '공실', '임대중', '성업', '전세', '렌트프리', '수익', '명도']


def extract_tenant(desc):
    if not desc:
        return ''
    hits = []
    # 보증금/월세 조건 패턴
    for m in re.finditer(r'(?:보증금\s*)?([\d,]+)\s*/\s*([\d,]+)', desc):
        hits.append(f"{m.group(1)}/{m.group(2)}")
    for m in re.finditer(r'월세\s*([\d,]+)\s*(?:만)?', desc):
        hits.append(f"월세 {m.group(1)}")
    for m in re.finditer(r'(?:수익률|수익율)\s*([\d.]+)\s*%?', desc):
        hits.append(f"수익률 {m.group(1)}%")
    kw = [k for k in ['세입자', '임차인', '임대중', '성업', '공실', '명도', '실입주'] if k in desc]
    hits += kw
    # 중복 제거 (순서 유지)
    out, seen = [], set()
    for h in hits:
        if h not in seen:
            seen.add(h); out.append(h)
    return ' · '.join(out)


def parse(t):
    r = {'raw': t}
    # 매매가
    m = re.search(r'매매\s*((?:\d+억)?\s*[\d,]*)', t)
    price_txt = m.group(1).strip() if m else ''
    r['price'] = won(price_txt)
    r['price_txt'] = price_txt
    # 전용면적
    m = re.search(r'전용([\d.]+)\)', t)
    r['jeon'] = float(m.group(1)) if m else None
    r['py'] = r['jeon'] / PY if r['jeon'] else None
    # 계약면적
    m = re.search(r'계약([\d.]+)㎡', t)
    r['gye'] = float(m.group(1)) if m else None
    # 층
    m = re.search(r'\)\s*(B?\d+|옥탑|고|중|저)\s*/\s*(\d+)층', t)
    r['floor'] = m.group(1) if m else ''
    r['total_floor'] = m.group(2) if m else ''
    # 유형
    m = re.search(r'(중소형사무실|일반상가|복합상가|단지내상가|근린상가|상가점포|사무실|아파트상가)', t)
    r['typ'] = m.group(1) if m else ''
    # 설명 (따옴표 안)
    m = re.search(r'"([^"]+)"', t)
    r['desc'] = m.group(1).strip() if m else ''
    r['tenant'] = extract_tenant(r['desc'])
    # 중개사 그룹
    g = re.search(r'중개사\s*(\d+)곳', t)
    r['ncnt'] = int(g.group(1)) if g else 1
    r['ppp'] = round(r['price'] / r['py']) if (r['price'] and r['py']) else None
    return r


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('input')
    ap.add_argument('xlsx', nargs='?')
    ap.add_argument('--title', default='네이버 부동산 매매 매물')
    a = ap.parse_args()
    data = json.load(open(a.input, encoding='utf-8'))
    rows = [parse(t) for t in data]
    rows = [r for r in rows if r['price']]

    # 중복 제거: 같은 층 + 같은 매매가 + 전용평수 차이 0.15평 이내 → 동일 매물로 간주
    dedup = []
    for r in rows:
        hit = None
        for d in dedup:
            if d['floor'] == r['floor'] and d['price'] == r['price']:
                if d['py'] is not None and r['py'] is not None and abs(d['py'] - r['py']) <= 0.15:
                    hit = d; break
                if d['py'] is None and r['py'] is None:
                    hit = d; break
        if hit:
            hit['dupe'] += 1
            if not hit['tenant'] and r['tenant']:
                hit['tenant'] = r['tenant']
            if not hit['desc'] and r['desc']:
                hit['desc'] = r['desc']
            continue
        r['dupe'] = 1
        dedup.append(r)

    dedup.sort(key=lambda r: (float(r['floor'].replace('B', '-')) if re.match(r'^B?\d+$', r['floor']) else 99,
                              -(r['price'] or 0)))

    print(f"# {a.title}\n")
    print(f"원본 {len(rows)}건 → 중복 제거 후 **{len(dedup)}건**\n")
    print("| # | 층 | 전용평수 | 평당 매매단가(만원) | 매매가(만원) | 매매가 | 임차 관련 내용 | 매물설명 | 유형 | 중복 |")
    print("|---|---|---|---|---|---|---|---|---|---|")
    for i, r in enumerate(dedup, 1):
        print("| {} | {} | {} | {} | {} | {} | {} | {} | {} | {} |".format(
            i, (r['floor'] + '층') if r['floor'] else '',
            f"{r['py']:,.2f}" if r['py'] else '',
            f"{r['ppp']:,}" if r['ppp'] else '',
            f"{r['price']:,}" if r['price'] else '',
            r['price_txt'],
            r['tenant'] or '-',
            r['desc'].replace('|', '/'),
            r['typ'], r['dupe']))

    ppps = sorted(r['ppp'] for r in dedup if r['ppp'])
    print("\n### 요약")
    print(f"- 총 {len(dedup)}건 (중복 {len(rows)-len(dedup)}건 제거)")
    if ppps:
        mid = ppps[len(ppps)//2]
        print(f"- 평당 매매단가: **{ppps[0]:,} ~ {ppps[-1]:,}만원/평** (중앙값 {mid:,})")
    tn = sum(1 for r in dedup if r['tenant'])
    print(f"- 임차 관련 내용 표기 매물: {tn}건")

    if a.xlsx:
        write_xlsx(dedup, a.xlsx, a.title)
        print(f"\n엑셀: {a.xlsx}")


def write_xlsx(rows, path, title):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    wb = Workbook(); ws = wb.active; ws.title = '매매매물'
    ws.merge_cells('A1:J1')
    ws['A1'] = title; ws['A1'].font = Font(bold=True, size=13)
    ws['A2'] = '단위: 만원 | 전용평수=전용㎡÷3.3058 | 평당단가=매매가÷전용평 | 동일 층·면적·금액 중복 1건으로 통합'
    ws['A2'].font = Font(size=9, color='666666')
    headers = ['#', '층', '전용평수', '전용㎡', '평당 매매단가(만원)', '매매가(만원)',
               '임차 관련 내용', '매물설명', '유형', '중복등록']
    hr = 4
    hf = PatternFill('solid', fgColor='2E7D32'); hfont = Font(bold=True, color='FFFFFF', size=10)
    thin = Side(style='thin', color='D0D0D0'); bd = Border(thin, thin, thin, thin)
    for j, h in enumerate(headers, 1):
        c = ws.cell(hr, j, h); c.fill = hf; c.font = hfont; c.border = bd
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    r = hr + 1
    for i, x in enumerate(rows, 1):
        vals = [i, (x['floor'] + '층') if x['floor'] else '',
                round(x['py'], 2) if x['py'] else None, x['jeon'],
                x['ppp'], x['price'], x['tenant'], x['desc'], x['typ'], x['dupe']]
        for j, v in enumerate(vals, 1):
            c = ws.cell(r, j, v); c.border = bd
            c.alignment = Alignment(horizontal='center' if j in (1, 2) else
                                    ('left' if j in (7, 8, 9) else 'right'),
                                    vertical='center', wrap_text=(j in (7, 8)))
            if j in (3, 4, 5, 6) and isinstance(v, (int, float)):
                c.number_format = '#,##0.00' if j in (3, 4) else '#,##0'
        if i % 2 == 0:
            for j in range(1, 11):
                ws.cell(r, j).fill = PatternFill('solid', fgColor='F5F9F5')
        r += 1
    for j, w in enumerate([4, 7, 10, 10, 18, 14, 30, 46, 14, 9], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = 'A5'
    ws.auto_filter.ref = f"A{hr}:J{r-1}"
    wb.save(path)


if __name__ == '__main__':
    main()
