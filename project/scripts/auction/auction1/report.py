"""수집 결과 → xlsx / md 출력.

지하 제외는 여기서 하지 않는다. 분류는 classify() 가 하고,
report 는 받은 그룹을 그대로 기록만 한다.
"""

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .parse import COLUMNS

__all__ = [
    "Classified",
    "classify",
    "write_xlsx",
    "write_markdown",
    "write_ranking_sheet",
    "RANKING_SHEET",
    "DEEPEN_SHEET",
    "RANKING_COLUMNS",
]

RANKING_SHEET = "단타랭킹"
DEEPEN_SHEET = "심화"  # auction1_deepen 이 쓰고 auction1_screen 이 읽는다

# 랭킹 시트 컬럼. 판단에 필요한 근거를 옆에 다 붙여 둔다 — 점수만 보고 믿지 말라는 뜻.
RANKING_COLUMNS = [
    "순위",
    "단타점수",
    "사건번호",
    "물건번호",
    "물건종류",
    "시군구",
    "소재지",
    "층",
    "호수",
    "건물면적평",
    "감정가",
    "최저가",
    "최저가율",
    "저감률",
    "평당가만원",
    "동네감정평당가",
    "비교단위",
    "비교단지수",
    "비교신뢰도",
    "저평가율",
    "유찰횟수",
    "조회수",
    "매각기일",
    "임차인유무",
    "보증금합계",
    "월세합계",
    "특이사항",
    "리스크감점",
    "리스크사유",
    "환금경고",
    "규제",
    "옥션원링크",
]

# 지하 판정 근거는 제외/미상 시트에서만 의미가 있다.
_REASON_COLUMN = "지하판정근거"

# md 표에 넣을 컬럼 (엑셀보다 좁게)
_MD_COLUMNS = [
    "사건번호",
    "물건종류",
    "시군구",
    "소재지",
    "층",
    "건물면적평",
    "감정가",
    "최저가",
    "최저가율",
    "유찰횟수",
    "매각기일",
]


@dataclass
class Classified:
    """지하 판정 결과에 따른 3분류."""

    ground: list[dict[str, Any]]  # 지상 — 검토 대상
    basement: list[dict[str, Any]]  # 지하 — 제외
    unknown: list[dict[str, Any]]  # 층 판정 불가 — 사람이 확인


def classify(rows: list[dict[str, Any]]) -> Classified:
    """지하여부(True/False/None) 로 나눈다. None 을 조용히 버리지 않는다."""
    result = Classified(ground=[], basement=[], unknown=[])
    for row in rows:
        match row.get("지하여부"):
            case True:
                result.basement.append(row)
            case False:
                result.ground.append(row)
            case _:
                result.unknown.append(row)
    return result


def _sort_key(row: dict[str, Any]):
    """매각기일 → 최저가 순. 값이 없는 건은 뒤로."""
    date = row.get("매각기일") or "9999-99-99"
    price = row.get("최저가")
    return (date, price if price is not None else float("inf"))


# --- 엑셀 --------------------------------------------------------------------

_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_MONEY_COLUMNS = {"감정가", "최저가", "보증금합계", "월세합계"}
_WIDTHS = {"소재지": 52, "특이사항": 26, "규제": 24, "상세링크": 20, "지하판정근거": 26}


def write_xlsx(path: Path, groups: Classified, meta: dict[str, str]) -> Path:
    """3개 시트(결과 / 제외_지하 / 층미상)로 저장."""
    workbook = Workbook()
    workbook.remove(workbook.active)

    _add_sheet(workbook, "결과", sorted(groups.ground, key=_sort_key), COLUMNS, meta)
    _add_sheet(
        workbook,
        "제외_지하",
        sorted(groups.basement, key=_sort_key),
        COLUMNS + [_REASON_COLUMN],
        meta,
    )
    _add_sheet(
        workbook,
        "층미상",
        sorted(groups.unknown, key=_sort_key),
        COLUMNS + [_REASON_COLUMN],
        meta,
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def _add_sheet(workbook, title, rows, columns, meta) -> None:
    sheet = workbook.create_sheet(title)

    sheet.append([f"{title} — {meta['조건']} · 수집 {meta['수집시각']} · {len(rows)}건"])
    sheet["A1"].font = Font(bold=True, size=11)
    sheet.append([])

    sheet.append(columns)
    header_row = sheet.max_row
    for index, _ in enumerate(columns, start=1):
        cell = sheet.cell(row=header_row, column=index)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        sheet.append([row.get(column) for column in columns])

    for index, column in enumerate(columns, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = _WIDTHS.get(column, 13)
        if column in _MONEY_COLUMNS:
            for cell in sheet[letter][header_row:]:
                cell.number_format = "#,##0"

    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
    sheet.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(len(columns))}{sheet.max_row}"
    )


# --- 랭킹 시트 (기존 파일에 탭 추가) ------------------------------------------

_TOP_FILL = PatternFill("solid", fgColor="FFF3D6")  # 상위권 강조
_LINK_FONT = Font(color="0563C1", underline="single")
_RANKING_WIDTHS = {
    "소재지": 52,
    "특이사항": 24,
    "리스크사유": 22,
    "환금경고": 20,
    "규제": 22,
    "비교신뢰도": 11,
    "옥션원링크": 12,
}
HIGHLIGHT_TOP = 20


def write_ranking_sheet(path: Path, scored: list[dict[str, Any]], meta: dict[str, str],
                        top: int | None = None) -> Path:
    """기존 수집 xlsx 에 단타 랭킹 탭을 추가한다 (이미 있으면 갈아끼운다).

    다른 시트(결과·제외_지하·층미상)는 건드리지 않는다.
    """
    from openpyxl import load_workbook

    rows = scored[:top] if top else scored
    workbook = load_workbook(path)
    if RANKING_SHEET in workbook.sheetnames:
        del workbook[RANKING_SHEET]
    sheet = workbook.create_sheet(RANKING_SHEET)

    sheet.append([
        f"{RANKING_SHEET} — {meta['조건']} · {len(rows)}건 (점수 내림차순) · 산출 {meta['수집시각']}"
    ])
    sheet["A1"].font = Font(bold=True, size=11)
    sheet.append(["점수는 볼 순서를 정할 뿐이다. 인근 실거래 마진·명도 난이도·인수 권리는 "
                  "목록에 없으므로 상세와 현장에서 확인해야 한다."])
    sheet["A2"].font = Font(size=9, color="808080")
    sheet.append([])

    sheet.append(RANKING_COLUMNS)
    header_row = sheet.max_row
    for index, _ in enumerate(RANKING_COLUMNS, start=1):
        cell = sheet.cell(row=header_row, column=index)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    link_index = RANKING_COLUMNS.index("옥션원링크") + 1
    for rank, row in enumerate(rows, start=1):
        values = [rank if column == "순위" else row.get(column) for column in RANKING_COLUMNS]
        values[link_index - 1] = None  # 링크 셀은 아래에서 하이퍼링크로 채운다
        sheet.append(values)

        excel_row = sheet.max_row
        url = row.get("상세링크")
        cell = sheet.cell(row=excel_row, column=link_index)
        if url:
            cell.value = "옥션원"
            cell.hyperlink = url
            cell.font = _LINK_FONT
        if rank <= HIGHLIGHT_TOP:
            for column_index in range(1, len(RANKING_COLUMNS) + 1):
                sheet.cell(row=excel_row, column=column_index).fill = _TOP_FILL

    for index, column in enumerate(RANKING_COLUMNS, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = _RANKING_WIDTHS.get(column, 12)
        if column in _MONEY_COLUMNS:
            for cell in sheet[letter][header_row:]:
                cell.number_format = "#,##0"

    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=4)
    sheet.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(len(RANKING_COLUMNS))}{sheet.max_row}"
    )
    workbook.save(path)
    return path


# --- 마크다운 ----------------------------------------------------------------


def write_markdown(path: Path, groups: Classified, meta: dict[str, str],
                   warnings: list[str]) -> Path:
    ground = sorted(groups.ground, key=_sort_key)
    lines = [
        f"# 옥션원 경매검색 — {meta['조건']}",
        "",
        f"> 수집: {meta['수집시각']} · 출처: auction1.co.kr 종합검색",
        "",
        "## 요약",
        "",
        "| 항목 | 건수 |",
        "|---|---|",
        f"| 사이트 표시 총건수 | {meta['총건수']} |",
        f"| 수집 건수 | {meta['수집건수']} |",
        f"| **검토 대상 (지상)** | **{len(ground)}** |",
        f"| 제외 (지하) | {len(groups.basement)} |",
        f"| 층미상 (확인 필요) | {len(groups.unknown)} |",
        "",
    ]

    if warnings:
        lines += ["## ⚠️ 경고", ""]
        lines += [f"- {warning}" for warning in warnings]
        lines.append("")

    lines += ["## 검토 대상", ""]
    if ground:
        lines += _md_table(ground)
    else:
        lines.append("조건에 맞는 지상 물건이 없다.")
    lines.append("")

    if groups.unknown:
        lines += [
            "## 층미상 — 사람이 확인 필요",
            "",
            "층 표기를 해석하지 못한 건이다. 조용히 버리지 않고 남긴다.",
            "",
        ]
        lines += _md_table(groups.unknown)
        lines.append("")

    lines += [
        "---",
        "",
        "본 자료는 옥션원 유료계정으로 조회한 개인 검토용 자료다. 외부 공유·재배포 금지.",
        "",
    ]

    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def _md_table(rows: list[dict[str, Any]]) -> list[str]:
    header = "| " + " | ".join(_MD_COLUMNS) + " |"
    divider = "|" + "---|" * len(_MD_COLUMNS)
    body = []
    for row in rows:
        cells = [_md_cell(column, row.get(column)) for column in _MD_COLUMNS]
        body.append("| " + " | ".join(cells) + " |")
    return [header, divider, *body]


def _md_cell(column: str, value: Any) -> str:
    if value is None or value == "":
        return "-"
    if column in _MONEY_COLUMNS:
        return _won(value)
    if column == "최저가율":
        return f"{value}%"
    if column == "유찰횟수":
        return "신건" if value == 0 else f"{value}회"
    return str(value).replace("|", "\\|")


def _won(amount: int) -> str:
    """216000000 → '2.16억'."""
    if amount >= 100_000_000:
        return f"{amount / 100_000_000:.2f}억"
    return f"{amount / 10_000:,.0f}만"


def build_meta(query_desc: str, total: int | None, collected: int) -> dict[str, str]:
    return {
        "조건": query_desc,
        "수집시각": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "총건수": str(total) if total is not None else "확인불가",
        "수집건수": str(collected),
    }
