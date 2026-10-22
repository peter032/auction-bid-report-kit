"""입찰보고서 PPTX 렌더링 부품 (1280×720 px 좌표계, 네이티브 편집가능 도형·표).

봉담 DS타워(2025타경1684) · 파크앤시티Ⅱ(2024타경78261) 빌더에서 추출.
"""

import os

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

PRIMARY = "C1001B"; PRIMARY_L = "E8344A"; PRIMARY_D = "8B0013"
DARK = "2D2D2D"; GRAY = "555555"; GRAY_L = "F5F6F8"; GRAY_M = "E2E2E2"; WHITE = "FFFFFF"
OK = "16A34A"; OK_D = "15803D"; WARN = "D97706"; RISK = "DC2626"; BLUE = "2563EB"
MUTE = "999999"; MUTE2 = "AAAAAA"
LEVEL = {"risk": RISK, "warn": WARN, "info": BLUE, "ok": OK, "primary": PRIMARY}
FONT = "맑은 고딕"
NO_STYLE = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'
ALIGN = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}
ANCH = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}


def P(px):
    return Emu(int(round(px * 9525)))


def F(px):
    return Pt(px * 0.75)


def C(h):
    return RGBColor.from_string(h)


def _fix_font(run):
    rPr = run._r.get_or_add_rPr()
    for tag in ('a:ea', 'a:cs'):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {})
            rPr.append(e)
        e.set('typeface', FONT)


def TB(s, x, y, w, h, valign='t'):
    sh = s.shapes.add_textbox(P(x), P(y), P(w), P(h))
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = ANCH[valign]
    return sh, tf


def runs_of(value):
    """문자열 또는 [[텍스트, {스타일}], ...] 또는 **굵게** 마크업 → run 목록."""
    if isinstance(value, list):
        return [(r, {}) if isinstance(r, str) else (r[0], r[1] if len(r) > 1 else {}) for r in value]
    parts, out = str(value).split("**"), []
    for i, part in enumerate(parts):
        if part:
            out.append((part, {'bold': True} if i % 2 else {}))
    return out


def PARA(tf, runs, size=14, bold=False, color=DARK, align='l', lh=1.35, space_before=0, space_after=0):
    ps = tf.paragraphs
    p = ps[0] if (len(ps) == 1 and not ps[0].runs) else tf.add_paragraph()
    p.alignment = ALIGN[align]
    p.line_spacing = lh
    p.space_before = F(space_before)
    p.space_after = F(space_after)
    for t, st in runs_of(runs):
        run = p.add_run()
        run.text = t
        f = run.font
        f.name = FONT
        f.size = F(st.get('size', size))
        f.bold = st.get('bold', bold)
        f.color.rgb = C(st.get('color', color))
        _fix_font(run)
    return p


def TEXT(s, x, y, w, h, runs, valign='t', **kw):
    sh, tf = TB(s, x, y, w, h, valign)
    PARA(tf, runs, **kw)
    return sh


def RECT(s, x, y, w, h, fill=None, line=None, lw=1, radius=None, shape=MSO_SHAPE.RECTANGLE):
    if radius is not None:
        shape = MSO_SHAPE.ROUNDED_RECTANGLE
    sh = s.shapes.add_shape(shape, P(x), P(y), P(w), P(h))
    if radius is not None:
        sh.adjustments[0] = min(0.5, radius / float(min(w, h)))
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = C(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = C(line)
        sh.line.width = Pt(lw * 0.75)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def GRAD(s, x, y, w, h, c1, c2):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, P(x), P(y), P(w), P(h))
    sh.fill.gradient()
    sh.fill.gradient_angle = 45
    st = sh.fill.gradient_stops
    st[0].color.rgb, st[0].position = C(c1), 0.0
    st[1].color.rgb, st[1].position = C(c2), 1.0
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def LINE(s, x1, y1, x2, y2, color=DARK, w=1.3):
    cn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, P(x1), P(y1), P(x2), P(y2))
    cn.line.color.rgb = C(color)
    cn.line.width = Pt(w * 0.75)
    return cn


def PIC(s, path, x, y, w, h):
    """박스 안에 비율 유지로 맞춘다. 파일이 없으면 회색 자리표시 + 파일명(조용한 누락 금지)."""
    if not path or not os.path.exists(path):
        RECT(s, x, y, w, h, fill=GRAY_L, line=GRAY_M)
        TEXT(s, x, y + h / 2 - 10, w, 20, f"이미지 없음: {os.path.basename(str(path))}", size=11, color=MUTE, align='c')
        return None
    iw, ih = Image.open(path).size
    sc = min(w / iw, h / ih)
    dw, dh = iw * sc, ih * sc
    dx, dy = x + (w - dw) / 2, y + (h - dh) / 2
    s.shapes.add_picture(path, P(dx), P(dy), P(dw), P(dh))
    return dx, dy, dw, dh, sc


def _set_border(cell, edges, color, w=1):
    tcPr = cell._tc.get_or_add_tcPr()
    order = ['a:lnL', 'a:lnR', 'a:lnT', 'a:lnB']
    for e in edges:
        tag = {'L': 'a:lnL', 'R': 'a:lnR', 'T': 'a:lnT', 'B': 'a:lnB'}[e]
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = tcPr.makeelement(qn(tag), {'w': str(int(w * 12700 * 0.75)), 'cap': 'flat', 'cmpd': 'sng', 'algn': 'ctr'})
        fill = ln.makeelement(qn('a:solidFill'), {})
        fill.append(fill.makeelement(qn('a:srgbClr'), {'val': color}))
        ln.append(fill)
        tcPr.insert(sum(1 for t in order[:order.index(tag)] if tcPr.find(qn(t)) is not None), ln)


def _cell(cell, runs, fill, size, bold, color, align):
    cell.fill.solid()
    cell.fill.fore_color.rgb = C(fill)
    cell.text_frame.word_wrap = True
    cell.margin_left = cell.margin_right = P(10)
    cell.margin_top = cell.margin_bottom = P(4)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    PARA(cell.text_frame, runs, size=size, bold=bold, color=color, align=align, lh=1.3)


def GRID(s, x, y, rows, colw, rowh=36, size=12.5, good=(), bad=(), left=(0,), head=DARK, header=True):
    """rows[0] = 헤더. good/bad = 강조 행 인덱스(1-base 데이터 행). header=False 면 머리글 없는 키-값 표."""
    if not header:
        rows = [None] + list(rows)
    body = rows if header else rows[1:]
    gf = s.shapes.add_table(len(body), len(colw), P(x), P(y), P(sum(colw)), P(rowh * len(body)))
    t = gf.table
    t.first_row = t.horz_banding = False
    tblPr = t._tbl.find(qn('a:tblPr'))
    sid = tblPr.find(qn('a:tableStyleId'))
    if sid is None:
        sid = tblPr.makeelement(qn('a:tableStyleId'), {})
        tblPr.append(sid)
    sid.text = NO_STYLE
    for i, cw in enumerate(colw):
        t.columns[i].width = P(cw)
    for i, row in enumerate(rows):
        if row is None:
            continue
        r = i if header else i - 1
        t.rows[r].height = P(rowh)
        for j, val in enumerate(row):
            c = t.cell(r, j)
            al = 'l' if j in left else 'c'
            val = "" if val is None else str(val)
            if i == 0:
                _cell(c, val, head, size - 0.5, True, WHITE if head == DARK else GRAY, al)
            elif i in bad:
                _cell(c, val, "FFF1F2", size, True, PRIMARY_D, al)
                _set_border(c, 'TB', PRIMARY, 2)
            elif i in good:
                _cell(c, val, "F0FDF4", size, True, OK_D, al)
                _set_border(c, 'TB', OK, 2)
            else:
                _cell(c, val, "FAFAFB" if i % 2 == 0 else WHITE, size, False, DARK, al)
                _set_border(c, 'B', GRAY_M)
    return t


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = P(1280)
        self.prs.slide_height = P(720)
        self.page = 0

    def slide(self, chapter=None, title=None, subtitle=None):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.page += 1
        if chapter:
            TEXT(s, 60, 32, 1160, 20, chapter, size=13, bold=True, color=PRIMARY, lh=1.0)
            TEXT(s, 60, 54, 1160, 42, title, size=28, bold=True, lh=1.15)
            TEXT(s, 60, 100, 1160, 22, subtitle or "", size=14.5, color=GRAY, lh=1.0)
            RECT(s, 60, 131, 52, 4, fill=PRIMARY)
            TEXT(s, 1150, 694, 70, 16, str(self.page), size=12, color=MUTE2, align='r', lh=1.0)
        return s

    def save(self, path):
        self.prs.save(path)


def foot(s, txt, x=60, y=690, w=1060, size=11):
    TEXT(s, x, y, w, 30, txt, size=size, color=MUTE, lh=1.3)


def KPI(s, x, y, w, h, label, value, unit=None, sub=None, red=False, vs=28):
    RECT(s, x, y, w, h, fill=("FFF1F2" if red else GRAY_L), radius=12)
    sh, tf = TB(s, x + 14, y + 10, w - 28, h - 20, 'm')
    PARA(tf, label, size=12.5, bold=True, color=GRAY, align='c', lh=1.2, space_after=4)
    col = PRIMARY if red else DARK
    runs = [(str(value), {'size': vs, 'bold': True, 'color': col})]
    if unit:
        runs.append((unit, {'size': 15, 'bold': True, 'color': col}))
    PARA(tf, runs, align='c', lh=1.1)
    if sub:
        PARA(tf, sub, size=11.5, color=GRAY, align='c', lh=1.25, space_before=4)


def CARD(s, x, y, w, h, color, title, body, fill=GRAY_L, tsize=14.5, bsize=12.5):
    RECT(s, x, y, w, h, fill=fill, radius=10)
    RECT(s, x, y, 4, h, fill=color)
    sh, tf = TB(s, x + 22, y + 12, w - 40, h - 24, 'm')
    PARA(tf, title, size=tsize, bold=True, color=color, lh=1.2, space_after=6)
    for b in (body if isinstance(body, list) else [body]):
        PARA(tf, b, size=bsize, lh=1.45, space_after=2)
